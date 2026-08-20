import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

def generate_report(output_path: str, title: str, categories: list, plan: list, fact: list):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws.views.sheetView[0].showGridLines = True

    # Заголовок
    ws["A1"] = title
    ws["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="1A365D")

    headers = ["Категория", "План, ₽", "Факт, ₽", "Отклонение, ₽", "% Выполнения"]
    ws.append([]); ws.append(headers)

    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    for col in range(1, 6):
        c = ws.cell(row=3, column=col)
        c.fill = header_fill
        c.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")

    thin = Border(left=Side(style='thin', color='CBD5E0'), right=Side(style='thin', color='CBD5E0'),
                  top=Side(style='thin', color='CBD5E0'), bottom=Side(style='thin', color='CBD5E0'))

    start = 4
    for i, (cat, p, f) in enumerate(zip(categories, plan, fact), start):
        ws.cell(row=i, column=1, value=cat).border = thin
        ws.cell(row=i, column=2, value=p).number_format = '#,##0'
        ws.cell(row=i, column=3, value=f).number_format = '#,##0'
        ws.cell(row=i, column=4, value=f"=C{i}-B{i}").number_format = '#,##0'
        ws.cell(row=i, column=5, value=f"=C{i}/B{i}").number_format = '0.0%'
        for c in range(1, 6): ws.cell(row=i, column=c).border = thin

    tot = start + len(categories)
    ws.cell(row=tot, column=1, value="ИТОГО").font = Font(name="Segoe UI", bold=True)
    ws.cell(row=tot, column=2, value=f"=SUM(B4:B{tot-1})").number_format = '#,##0'
    ws.cell(row=tot, column=3, value=f"=SUM(C4:C{tot-1})").number_format = '#,##0'
    ws.cell(row=tot, column=4, value=f"=C{tot}-B{tot}").number_format = '#,##0'
    ws.cell(row=tot, column=5, value=f"=C{tot}/B{tot}").number_format = '0.0%'
    for c in range(1, 6): 
        ws.cell(row=tot, column=c).border = thin

    # Диаграмма
    chart = BarChart()
    chart.title = "План vs Факт"
    chart.width = 14
    chart.height = 8
    chart.add_data(Reference(ws, min_col=2, min_row=3, max_col=3, max_row=tot-1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=tot-1))
    ws.add_chart(chart, "G3")

    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = 20

    wb.save(output_path)
    print(f"Отчет сохранен в {output_path}")

if __name__ == "__main__":
    generate_report("test_report.xlsx", "Исполнение бюджета", ["Металл", "Логистика", "IT"], [5000000, 1200000, 400000], [4800000, 1350000, 390000])
