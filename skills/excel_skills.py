"""Excel/CSV analysis helpers used by OpenClaw."""
import json
import os
import re

import duckdb
import openpyxl
import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from skills.sql_security import validate_read_only_query


def _table_name(sheet_name: str) -> str:
    name = re.sub(r"[^a-zA-Z0-9_]+", "_", str(sheet_name).strip().lower()).strip("_")
    return name or "sheet"


def _read_tables(file_path: str) -> list[tuple[str, pd.DataFrame]]:
    suffix = os.path.splitext(file_path)[1].lower()
    if suffix == ".csv":
        return [("data", pd.read_csv(file_path))]
    excel_file = pd.ExcelFile(file_path)
    return [(_table_name(sheet), pd.read_excel(file_path, sheet_name=sheet)) for sheet in excel_file.sheet_names]


class ExcelSkillsHandler:
    @staticmethod
    def excel_sql_query(file_path: str, query: str) -> str:
        if not os.path.exists(file_path):
            return f"Ошибка: Файл {file_path} не найден."
        try:
            safe_query = validate_read_only_query(query)
            conn = duckdb.connect(database=":memory:")
            # Lock file/network access at the engine level; regex validation above
            # is only the first layer and must not be the only boundary.
            conn.execute("SET enable_external_access=false")
            try:
                for table_name, df in _read_tables(file_path):
                    df.columns = [re.sub(r"[^a-zA-Z0-9_]+", "_", str(c).strip().lower()).strip("_") or "column" for c in df.columns]
                    conn.register(table_name, df)
                result = conn.execute(safe_query).df()
            finally:
                conn.close()
            if result.empty:
                return "Запрос выполнен успешно, записей не найдено."
            return result.to_markdown(index=False)
        except Exception as err:
            return f"Ошибка выполнения SQL: {str(err)}"

    @staticmethod
    def excel_table_search(file_path: str, keyword: str) -> str:
        if not os.path.exists(file_path):
            return f"Ошибка: Файл {file_path} не найден."
        try:
            matches = []
            kw = str(keyword).lower()
            for sheet_name, df in _read_tables(file_path):
                df = df.dropna(how="all").fillna("")
                for idx, row in df.iterrows():
                    row_dict = row.to_dict()
                    row_str = " | ".join(f"{k}: {v}" for k, v in row_dict.items() if str(v).strip())
                    if kw in row_str.lower():
                        matches.append({"sheet": sheet_name, "row": int(idx) + 2, "data": row_dict})
            if not matches:
                return f"По запросу '{keyword}' совпадений в таблице не найдено."
            return json.dumps(matches[:15], ensure_ascii=False, indent=2, default=str)
        except Exception as err:
            return f"Ошибка поиска по таблице: {str(err)}"

    @staticmethod
    def excel_generate_report(output_path: str, title: str, categories: list, plan_values: list, fact_values: list) -> str:
        try:
            if not (len(categories) == len(plan_values) == len(fact_values)):
                return "Ошибка: категории, план и факт должны иметь одинаковую длину."
            if not categories:
                return "Ошибка: отчёт не может быть пустым."
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Сводка"
            ws["A1"] = title
            ws["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="1A365D")
            ws.append([])
            ws.append(["Категория", "План, ₽", "Факт, ₽", "Отклонение, ₽", "% Выполнения"])
            fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
            for col in range(1, 6):
                cell = ws.cell(row=3, column=col)
                cell.fill = fill
                cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            border = Border(*(Side(style="thin", color="CBD5E0") for _ in range(4)))
            for row, (cat, plan, fact) in enumerate(zip(categories, plan_values, fact_values), 4):
                ws.cell(row=row, column=1, value=str(cat))
                ws.cell(row=row, column=2, value=plan).number_format = "#,##0"
                ws.cell(row=row, column=3, value=fact).number_format = "#,##0"
                ws.cell(row=row, column=4, value=f"=C{row}-B{row}").number_format = "#,##0"
                ws.cell(row=row, column=5, value=f'=IFERROR(C{row}/B{row},0)').number_format = "0.0%"
                for col in range(1, 6):
                    ws.cell(row=row, column=col).border = border
            total = 4 + len(categories)
            ws.cell(row=total, column=1, value="ИТОГО").font = Font(name="Segoe UI", bold=True)
            ws.cell(row=total, column=2, value=f"=SUM(B4:B{total-1})").number_format = "#,##0"
            ws.cell(row=total, column=3, value=f"=SUM(C4:C{total-1})").number_format = "#,##0"
            ws.cell(row=total, column=4, value=f"=C{total}-B{total}").number_format = "#,##0"
            ws.cell(row=total, column=5, value=f'=IFERROR(C{total}/B{total},0)').number_format = "0.0%"
            for col in range(1, 6):
                ws.cell(row=total, column=col).border = border
            chart = BarChart()
            chart.title = "План vs Факт"
            chart.add_data(Reference(ws, min_col=2, min_row=3, max_col=3, max_row=total - 1), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=total - 1))
            ws.add_chart(chart, "G3")
            for col in ws.columns:
                ws.column_dimensions[get_column_letter(col[0].column)].width = 20
            wb.save(output_path)
            return f"Отчёт успешно сгенерирован и сохранён: {output_path}"
        except Exception as err:
            return f"Ошибка генерации отчёта: {str(err)}"

    def execute_tool(self, tool_name: str, arguments: dict) -> str:
        method = getattr(self, tool_name, None)
        if not method:
            return f"Инструмент '{tool_name}' не зарегистрирован."
        return method(**arguments)
